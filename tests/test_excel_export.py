"""The exports have to be a workbook, not a screenshot of one.

Every check here opens the finished .xlsx with openpyxl and looks at real
cells. The thing that makes an export unprofessional is rarely that it is
missing data — it is that money arrives as text, so it will not sum; that
the header scrolls away on the second screen; that columns are too narrow
to show a figure; and that the file does not say which company or which
month it belongs to once it is an attachment in somebody's inbox.

It also pins the bug the old CSV had, which was not cosmetic at all: it
joined each worker to each of their attendance rows, so pay policy was
repeated once per marked day and any worker with no attendance was absent
from the file entirely.
"""
import io
import os
import re
import sys
import urllib.parse
import urllib.request
import http.cookiejar
import subprocess
import time

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')

from openpyxl import load_workbook

from app import app
import models as m
from timeutil import today_ist

PORT = 5121
B = 'http://127.0.0.1:%d' % PORT
fails = []


def check(label, got, want):
    ok = got == want
    if not ok:
        fails.append((label, got, want))
    print('  %-5s %-50s %r%s' % ('OK' if ok else 'FAIL', label, got,
                                 '' if ok else '  (want %r)' % (want,)))


def truthy(label, got):
    ok = bool(got)
    if not ok:
        fails.append((label, got, 'something truthy'))
    print('  %-5s %-50s %r' % ('OK' if ok else 'FAIL', label, got))


with app.app_context():
    # Outside a request there is no tenant filter, so the demo account's
    # own workers are picked out by hand — another company's worker is
    # correctly absent from the export and must not be looked for.
    demo = m.User.query.filter_by(username='demo').first()
    tenant_workers = m.Worker.query.filter_by(account_id=demo.account_id).all()
    TENANT_NAMES = {w.full_name for w in tenant_workers}
    WORKERS = len(tenant_workers)
    RECORDS = m.AttendanceRecord.query.count()
    NO_ATTENDANCE = [w.full_name for w in tenant_workers
                     if not w.attendance_records]

server = subprocess.Popen(
    [sys.executable, '-c',
     'import sys; sys.path.insert(0, %r);'
     'from app import app; import routes, bulk;'
     'app.run(host="127.0.0.1", port=%d)' % (ROOT, PORT)],
    cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
try:
    for _ in range(40):
        try:
            urllib.request.urlopen(B + '/login', timeout=1)
            break
        except Exception:
            time.sleep(0.5)
    else:
        sys.exit('could not start a server on %s' % B)

    cj = http.cookiejar.CookieJar()
    op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
    html = op.open(B + '/login').read().decode()
    tok = re.search(r'name="csrf_token"[^>]*value="([^"]+)"', html).group(1)
    op.open(B + '/login', urllib.parse.urlencode(
        {'username': 'demo@acme.test', 'password': 'demo12345',
         'csrf_token': tok}).encode()).read()

    def fetch(path):
        r = op.open(B + path)
        return r.headers, r.read()

    # ------------------------------------------------------------------
    print('--- payroll, as a workbook ---')
    today = today_ist()
    headers, blob = fetch('/payroll/export.xlsx?month=%d&year=%d'
                          % (today.month, today.year))
    check('served as a spreadsheet, not as text',
          'spreadsheetml' in headers['Content-Type'], True)
    truthy('filename says what it is',
           '.xlsx' in headers.get('Content-Disposition', ''))

    wb = load_workbook(io.BytesIO(blob))
    ws = wb['Payroll']
    check('one sheet, named', wb.sheetnames, ['Payroll'])

    # The title block: an attachment has to say whose and which month.
    head = '\n'.join(str(ws.cell(row=r, column=1).value or '')
                     for r in range(1, 4))
    truthy('names the company', len(head.splitlines()[0]) > 0)
    truthy('names the month', str(today.year) in head)
    truthy('says when it was generated', 'Generated' in head)

    # Row 5 is the header; the body starts at 6.
    headings = [c.value for c in ws[5] if c.value]
    truthy('has a header row', len(headings) > 15)
    check('money columns are present',
          all(h in headings for h in
              ('Gross Pay', 'Deductions', 'Net Payable')), True)
    check('the header is bold', ws.cell(row=5, column=1).font.bold, True)
    check('and frozen, so it survives scrolling', ws.freeze_panes, 'A6')
    truthy('a filter is on, for narrowing to one department',
           ws.auto_filter.ref)

    net_col = headings.index('Net Payable') + 1
    body = [ws.cell(row=r, column=net_col)
            for r in range(6, ws.max_row + 1)
            if ws.cell(row=r, column=2).value
            and not str(ws.cell(row=r, column=2).value).startswith('TOTAL')]
    truthy('there are worker rows', len(body))
    if body:
        # The whole point: a column of text that looks like money will not
        # sum and will not sort.
        check('net pay is a number, not text',
              all(isinstance(c.value, (int, float)) for c in body), True)
        check('and carries a rupee format',
              '₹' in (body[0].number_format or ''), True)
        widths = [ws.column_dimensions[k].width
                  for k in ('A', 'B', 'H', 'J')]
        check('columns are wide enough to show their contents',
              all(w and w >= 10 for w in widths), True)

    # The totals row, which the CSV wrote as a string padded with eighteen
    # empty cells to reach the right column.
    last = ws.max_row
    label = str(ws.cell(row=last, column=2).value or '')
    truthy('there is a totals row', label.startswith('TOTAL'))
    total_cell = ws.cell(row=last, column=net_col)
    check('the total is a number too',
          isinstance(total_cell.value, (int, float)), True)
    check('it is bold', total_cell.font.bold, True)
    if body and isinstance(total_cell.value, (int, float)):
        summed = round(sum(c.value or 0 for c in body), 2)
        check('and it equals the rows above it',
              round(total_cell.value, 2), summed)

    # ------------------------------------------------------------------
    print('\n--- the full export, and the join it used to do ---')
    headers, blob = fetch('/export_data.xlsx')
    wb = load_workbook(io.BytesIO(blob))
    check('one sheet per kind of thing', wb.sheetnames,
          ['Workers', 'Attendance'])

    ws = wb['Workers']
    names = [ws.cell(row=r, column=2).value
             for r in range(6, ws.max_row + 1)
             if ws.cell(row=r, column=2).value]
    # By employee ID and not by name: two people in the demo data are both
    # called Arvind Kejriwal, which is ordinary and not a duplicated row.
    ids = [ws.cell(row=r, column=1).value
           for r in range(6, ws.max_row + 1)
           if ws.cell(row=r, column=1).value]
    check('every worker appears exactly once', len(ids), len(set(ids)))
    check('and all of them are there', len(ids), WORKERS)
    # The old CSV looped over each worker's attendance rows, so a worker
    # with none of them never reached the file at all. Those are exactly
    # the workers worth checking for.
    listed = set(names)
    absent = [n for n in NO_ATTENDANCE if n in TENANT_NAMES and n not in listed]
    print('        workers with no attendance at all:',
          [n for n in NO_ATTENDANCE if n in TENANT_NAMES] or 'none')
    check('and none of them is missing from the sheet', absent, [])

    # Said once each, not once per attendance row.
    count_col = [c.value for c in ws[5]].index('Attendance Records') + 1
    busiest = max((ws.cell(row=r, column=count_col).value or 0)
                  for r in range(6, ws.max_row + 1))
    print('        busiest worker has %d attendance rows; the old file '
          'repeated their pay rules that many times' % busiest)

    ws_a = wb['Attendance']
    att_rows = sum(1 for r in range(6, ws_a.max_row + 1)
                   if ws_a.cell(row=r, column=2).value)
    truthy('attendance has its own sheet with rows', att_rows)
    check('each attendance row is listed once',
          att_rows <= RECORDS, True)
    check('both sheets freeze their header',
          (ws.freeze_panes, ws_a.freeze_panes), ('A6', 'A6'))

    # ------------------------------------------------------------------
    print('\n--- the CSV endpoints still answer ---')
    for path in ('/payroll/export.csv', '/export_data'):
        hdrs, data = fetch(path)
        check('%s still works' % path, 'csv' in hdrs['Content-Type'], True)
finally:
    server.terminate()
    server.wait(timeout=10)

print()
if fails:
    print('FAILED %d of the checks above:' % len(fails))
    for label, got, want in fails:
        print('  %s: got %r, want %r' % (label, got, want))
    sys.exit(1)
print('The exports are workbooks.')
